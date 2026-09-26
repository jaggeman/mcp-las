import asyncio
import base64
import io
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock

import pytest
from openpyxl import load_workbook

from src import server
from src.mcp_tools import tools


def test_invalid_mcp_keys_do_not_create_rate_buckets(monkeypatch):
    limiter = Mock(return_value=True)
    monkeypatch.setattr(server.auth_service, 'validate_key', lambda key: None)
    monkeypatch.setattr(server.auth_service, 'check_rate_limit', limiter)
    for n in range(305):
        assert server._check_rate_limit(f'invalid-{n}')['status'] == 'unauthorized'
    assert not limiter.called


def test_valid_and_anonymous_mcp_limits(monkeypatch):
    limiter = Mock(return_value=True)
    monkeypatch.setattr(server.auth_service, 'validate_key', lambda key: {'is_active': True})
    monkeypatch.setattr(server.auth_service, 'check_rate_limit', limiter)
    assert server._check_rate_limit('valid') is None
    assert limiter.call_args.kwargs['max_requests'] == 300
    assert server._check_rate_limit() is None
    assert limiter.call_args.kwargs['max_requests'] == 60


def test_excel_strings_are_not_formulas_and_filename_is_safe():
    result = tools.generate_turordningslista_excel(company_name='Bad\r\n"/公司', employees=[{
        'name': '=1+1', 'title': '=2+2', 'driftsenhet': '=3+3',
        'avtalsomrade': '=4+4', 'birth_date': '1990-01-01', 'notes': '=6+6',
    }])
    book = load_workbook(io.BytesIO(base64.b64decode(result['file_base64'])))
    for row in book.active.iter_rows(min_row=5):
        for cell in row:
            if cell.column != 7:
                assert cell.data_type != 'f'
    assert book.active['B5'].value.lstrip("'") == '=1+1'
    assert book.active['G5'].data_type == 'f'
    assert len(result['file_id']) >= 43
    assert result['file_name'].isascii()
    assert not any(c in result['file_name'] for c in '\r\n"/')
    response = asyncio.run(server.download_turordning_excel(SimpleNamespace(
        method='GET', query_params={'id': result['file_id']})))
    assert response.status_code == 200
    assert response.headers['cache-control'] == 'no-store'
    assert response.headers['referrer-policy'] == 'no-referrer'


def test_excel_store_expiration_capacity_and_size(monkeypatch):
    from src.services.download_store import DownloadStore
    store = DownloadStore(max_files=2, max_bytes=6, max_file_bytes=4, ttl=60)
    try:
        store['a'] = {'bytes': b'123', 'file_name': 'a.xlsx'}
        store['b'] = {'bytes': b'456', 'file_name': 'b.xlsx'}
        store['c'] = {'bytes': b'789', 'file_name': 'c.xlsx'}
        assert len(store) == 2 and 'a' not in store
        with pytest.raises(ValueError):
            store['large'] = {'bytes': b'12345'}
        monkeypatch.setattr('src.services.download_store.time.perf_counter', lambda: float('inf'))
        assert store.get('b') is None
        assert len(store) == 0
    finally:
        store.clear()


def test_excel_rejects_large_inputs():
    assert tools.generate_turordningslista_excel(employees=[{}] * 1001)['success'] is False
    assert tools.generate_turordningslista_excel(employees=[{'name': 'a' * 2001}])['success'] is False


def test_download_expiry_returns_404(monkeypatch):
    result = tools.generate_turordningslista_excel()
    monkeypatch.setattr('src.services.download_store.time.perf_counter', lambda: float('inf'))
    response = asyncio.run(server.download_turordning_excel(SimpleNamespace(
        method='GET', query_params={'id': result['file_id']})))
    assert response.status_code == 404


def test_store_cleans_up_without_further_requests():
    from threading import Event
    from src.services.download_store import DownloadStore
    cleaned = Event()
    store = DownloadStore(ttl=0.02)
    original = store._expire_in_background
    def cleanup():
        original()
        cleaned.set()
    store._expire_in_background = cleanup
    store['test'] = {'bytes': b'private'}
    try:
        assert cleaned.wait(timeout=2)
        assert not store._items  # Inspect without triggering lazy cleanup.
        assert store._bytes == 0
    finally:
        store.clear()


def test_store_total_bytes_are_bounded():
    from src.services.download_store import DownloadStore
    store = DownloadStore(max_files=10, max_bytes=6, max_file_bytes=4)
    try:
        store['a'] = {'bytes': b'1234'}
        store['b'] = {'bytes': b'5678'}
        assert list(store) == ['b']
        assert store._bytes == 4
    finally:
        store.clear()


def test_unknown_download_is_not_found():
    response = asyncio.run(server.download_turordning_excel(SimpleNamespace(
        method='GET', query_params={'id': '12345678'})))
    assert response.status_code == 404


def test_docker_excludes_environment_secrets():
    patterns = Path('.dockerignore').read_text().splitlines()
    assert '**/.env' in patterns
    assert '**/.env.*' in patterns
    assert '**/*credentials*.json' in patterns


def test_deploy_uses_a_dedicated_runtime_identity_and_locked_dependencies():
    workflow = Path('.github/workflows/ci.yml').read_text(encoding='utf-8')
    dockerfile = Path('Dockerfile').read_text(encoding='utf-8')
    assert '--service-account=mcp-las-runtime@paygap-prod.iam.gserviceaccount.com' in workflow
    assert 'pip-audit' in workflow
    assert '--require-hashes -r requirements.lock' in workflow
    assert '--require-hashes -r requirements.lock' in dockerfile


def test_repository_security_automation_is_enabled_and_actions_are_sha_pinned():
    workflows = list(Path('.github/workflows').glob('*.yml'))
    assert Path('.github/dependabot.yml').exists()
    assert Path('.github/workflows/codeql.yml').exists()
    for workflow in workflows:
        for line in workflow.read_text(encoding='utf-8').splitlines():
            if 'uses:' not in line:
                continue
            reference = line.split('uses:', 1)[1].strip().split()[0]
            assert '@' in reference
            revision = reference.rsplit('@', 1)[1]
            assert len(revision) == 40 and all(char in '0123456789abcdef' for char in revision)


def test_hosting_declares_browser_security_headers():
    config = Path('firebase.json').read_text(encoding='utf-8')
    assert 'Content-Security-Policy' in config
    assert 'Permissions-Policy' in config
