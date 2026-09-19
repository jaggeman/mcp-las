# -*- coding: utf-8 -*-
"""MCP-LAS tar inte emot personuppgifter — det är en lagtextserver.

Servern är publik (`--allow-unauthenticated`, kvot men ingen inloggning) och
delar Firestore-databas med Novro. Tre verktyg tog tidigare emot uppgifter om
namngivna personer: `generate_turordningslista_excel` (upp till 1000
anställda med födelsedatum och markering för vem som riskerar uppsägning),
`calculate_redundancy_turnorder_and_exceptions` (`employees_list`) och
`get_hr_document_template` (namn och personnummer).

Beräkningarna bakom dem är rena sorteringar och textmallar som den anropande
AI:n kan göra lokalt på data den redan har. Det enda servern tillförde var
filformatering — till priset av att persondata skickades till en publik
tjänst. Parametrarna är därför borttagna helt, och de här testerna låser fast
att de inte smyger tillbaka.
"""
import base64
import io
import inspect

import pytest
from openpyxl import load_workbook

from src.mcp_tools import tools


FORBIDDEN_PARAMETERS = [
    ("generate_turordningslista_excel", "employees"),
    ("calculate_redundancy_turnorder_and_exceptions", "employees_list"),
    ("get_hr_document_template", "employee_name"),
    ("get_hr_document_template", "personal_identity_number"),
]


@pytest.mark.parametrize("tool_name,parameter", FORBIDDEN_PARAMETERS)
def test_tool_has_no_personal_data_parameter(tool_name, parameter):
    signature = inspect.signature(getattr(tools, tool_name))
    assert parameter not in signature.parameters, (
        f"{tool_name} tar åter emot personuppgifter via '{parameter}'"
    )


@pytest.mark.parametrize("tool_name,parameter", FORBIDDEN_PARAMETERS)
def test_tool_rejects_the_personal_data_parameter_at_call_time(tool_name, parameter):
    # En anropare som fortfarande skickar det gamla argumentet ska få ett hårt
    # fel, inte tyst ignorering - annars tror den att uppgifterna användes.
    # Obligatoriska argument fylls i, annars skulle testet kunna passera på ett
    # TypeError om ETT ANNAT saknat argument och inte bevisa någonting.
    required = {"get_hr_document_template": {"template_type": "uppsagningsbesked_arbetsbrist"}}
    kwargs = dict(required.get(tool_name, {}))
    kwargs[parameter] = "Anna Andersson"
    with pytest.raises(TypeError, match=parameter):
        getattr(tools, tool_name)(**kwargs)


def _sheet1(result):
    workbook = load_workbook(io.BytesIO(base64.b64decode(result["file_base64"])))
    return workbook["Turordningslista"]


def test_the_workbook_is_an_empty_template():
    result = generate = tools.generate_turordningslista_excel(row_count=5)
    assert generate["success"] is True
    sheet = _sheet1(result)
    # Rad 4 är kolumnrubriker; datarader börjar på 5.
    for row_index in range(5, 10):
        for column in (2, 3, 4, 5, 6, 9, 12, 13):
            assert sheet.cell(row=row_index, column=column).value in (None, ""), (
                f"mallen är inte tom - rad {row_index} kolumn {column} har innehåll"
            )


def test_the_template_keeps_its_column_headers():
    sheet = _sheet1(tools.generate_turordningslista_excel(row_count=3))
    assert sheet.cell(row=4, column=2).value == "Namn"
    assert sheet.cell(row=4, column=6).value == "Anställningsdatum"


def test_each_blank_row_carries_a_datedif_formula_that_tolerates_an_empty_date():
    # Formeln är hela poängen med att leverera en mall: HR fyller i ett datum
    # i kolumn F och får anställningsdagarna automatiskt. Den måste därför
    # klara en tom cell utan att visa #VALUE! i varje oanvänd rad.
    sheet = _sheet1(tools.generate_turordningslista_excel(row_count=3))
    formula = sheet.cell(row=5, column=7).value
    assert "DATEDIF(F5" in str(formula), f"DATEDIF-formeln saknas: {formula!r}"
    assert str(formula).startswith("=IF(F5="), (
        f"formeln hanterar inte en tom datumcell: {formula!r}"
    )


def test_row_count_controls_how_many_rows_the_template_has():
    sheet = _sheet1(tools.generate_turordningslista_excel(row_count=7))
    assert sheet.cell(row=11, column=1).value == "EMP-007"
    assert sheet.cell(row=12, column=1).value is None


def test_row_count_is_bounded():
    assert tools.generate_turordningslista_excel(row_count=0)["success"] is False
    assert tools.generate_turordningslista_excel(row_count=1001)["success"] is False


def test_the_legal_rules_sheet_is_still_there():
    # Flik 2 är det verkliga värdet i verktyget: LAS 22 § och avtalens
    # undantagsregler. Den innehåller inga personuppgifter och blir kvar.
    workbook = load_workbook(io.BytesIO(base64.b64decode(
        tools.generate_turordningslista_excel(row_count=2)["file_base64"]
    )))
    assert "Undantagsregler & Lagstöd" in workbook.sheetnames


def test_the_hr_template_always_emits_placeholders():
    result = tools.get_hr_document_template(template_type="uppsagningsbesked_arbetsbrist")
    document = result["document_text"] if "document_text" in result else str(result)
    assert "[Arbetstagarens Förnamn Efternamn]" in document
    assert "[ÅÅÅÅMMDD-XXXX]" in document
